"""
Coffee data refresh (GitHub Actions).

  Futures (TradingView feed): ICE arabica KC, ICE robusta RC, B3 arabica ICF, FX   -> coffee/data/prices_cont.csv, prices_curve.csv
  Brazil physical (Notícias Agrícolas): CEPEA/ESALQ arabica & conillon indicators,
      co-op arabica prices (tipo 6 duro), conillon Vitória                         -> coffee/data/br_physical.csv
  Colombia (FNC): internal price daily, ex-dock monthly, ICO group indicators,
      production and exports monthly                                               -> coffee/data/col_*.csv
  Vietnam (giacaphe.com): Central Highlands average farm-gate robusta, VND/kg      -> coffee/data/vn_physical.csv
  ICO monthly report data (BMR-Com/coffee-analytics)                               -> coffee/data/ico_*.csv
  USDA PSD coffee                                                                  -> coffee/data/psd_coffee.csv
  CFTC COT coffee (BMR-Com/cftc-cot-full)                                          -> coffee/data/cot_coffee.csv

Usage: python coffee/scripts/fetch.py [--full] [--only a,b]
"""
import io, os, re, sys, time, zipfile, argparse, traceback
from datetime import date, timedelta, datetime
import requests
import pandas as pd

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
DATA = os.path.join(ROOT, "data")
os.makedirs(DATA, exist_ok=True)
H = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124 Safari/537.36"}
S = requests.Session(); S.headers.update(H)


def log(*a):
    print(time.strftime("%H:%M:%S"), *a, flush=True)


def get(u, **kw):
    for k in range(3):
        try:
            r = S.get(u, timeout=kw.pop("timeout", 60), **kw)
            return r
        except Exception as e:
            log("retry", u, e); time.sleep(2 + 3 * k)
    return None


def merge_save(new, fn, keys):
    p = os.path.join(DATA, fn)
    if os.path.exists(p):
        old = pd.read_csv(p)
        for k in keys:
            if k == "date":
                old[k] = pd.to_datetime(old[k]).dt.strftime("%Y-%m-%d")
        new = pd.concat([old, new])
    new = new.drop_duplicates(keys, keep="last").sort_values(keys)
    new.to_csv(p, index=False)
    log("saved", fn, new.shape)
    return new


# ---------------------------------------------------------------- futures
def _hist(tv, sym, ex, n=5000):
    from tvDatafeed import Interval
    for k in range(4):
        try:
            d = tv.get_hist(symbol=sym, exchange=ex, interval=Interval.in_daily, n_bars=n)
            if d is not None and len(d):
                d = d.reset_index()
                d["date"] = pd.to_datetime(d["datetime"]).dt.strftime("%Y-%m-%d")
                return d[["date", "open", "high", "low", "close", "volume"]]
        except Exception as e:
            log("tv err", sym, e)
        time.sleep(2 + 3 * k)
    return None


CONT = [("KC1!", "ICEUS"), ("KC2!", "ICEUS"), ("RC1!", "ICEEUR"), ("RC2!", "ICEEUR"), ("ICF1!", "BMFBOVESPA"),
        ("USDBRL", "FX_IDC"), ("USDCOP", "FX_IDC"), ("USDVND", "FX_IDC")]


def fetch_futures(full=False):
    from tvDatafeed import TvDatafeed
    tv = TvDatafeed()
    rows = []
    for sym, ex in CONT:
        d = _hist(tv, sym, ex)
        log("tv", sym, None if d is None else len(d))
        if d is not None:
            d.insert(0, "symbol", sym); rows.append(d)
    if rows:
        merge_save(pd.concat(rows), "prices_cont.csv", ["symbol", "date"])
    # live curve: next contracts of each market
    y = date.today().year
    want = [(f"KC{m}{yy}", "ICEUS") for yy in (y, y + 1, y + 2) for m in "HKNUZ"] + \
           [(f"RC{m}{yy}", "ICEEUR") for yy in (y, y + 1, y + 2) for m in "FHKNUX"]
    rows = []
    for sym, ex in want:
        d = _hist(tv, sym, ex, 400)
        if d is None:
            continue
        d.insert(0, "symbol", sym); rows.append(d)
    if rows:
        c = pd.concat(rows)
        c = c[pd.to_datetime(c.date) >= pd.Timestamp(date.today() - timedelta(days=400))]
        c.to_csv(os.path.join(DATA, "prices_curve.csv"), index=False)
        log("curve", c.symbol.nunique())


# ---------------------------------------------------------------- Brazil physical (Notícias Agrícolas)
NA = "https://www.noticiasagricolas.com.br/cotacoes/cafe/"
NA_SERIES = {"indicador-cepea-esalq-cafe-arabica": "cepea_arabica", "indicador-cepea-esalq-cafe-conillon": "cepea_conillon",
             "cafe-arabica-mercado-fisico-tipo-6-duro": "coop_arabica", "cafe-conillon-disponivel-vitoria-es": "vitoria_conillon"}


def _num(v):
    s = str(v).strip().replace("R$", "").strip()
    if s in ("", "-", "***", "nan"):
        return None
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except Exception:
        return None


def parse_na(html, series):
    """Each page has up to 15 day-blocks: a date header followed by a table."""
    out = []
    # split by the 'Fechamento: dd/mm/yyyy' style headers that precede each table
    parts = re.split(r"(\d{2}/\d{2}/\d{4})", html)
    tabs = []
    try:
        tabs = pd.read_html(io.StringIO(html), thousands=".", decimal=",")
    except Exception:
        return out
    dates = re.findall(r'class="fechamento"[^>]*>[^0-9]*(\d{2}/\d{2}/\d{4})', html) or \
            re.findall(r"Fechamento:?\s*(?:<[^>]+>)*\s*(\d{2}/\d{2}/\d{4})", html)
    for k, t in enumerate(tabs):
        cols = [str(c) for c in t.columns]
        if series.startswith("cepea"):
            if "Data" not in cols:
                continue
            for _, r in t.iterrows():
                v = _num(r[cols[1]])
                try:
                    d = datetime.strptime(str(r["Data"]), "%d/%m/%Y").strftime("%Y-%m-%d")
                except Exception:
                    continue
                if v:
                    out.append({"date": d, "series": series, "place": "CEPEA/ESALQ", "brl_bag": v})
        else:
            if k >= len(dates):
                continue
            d = datetime.strptime(dates[k], "%d/%m/%Y").strftime("%Y-%m-%d")
            for _, r in t.iterrows():
                v = _num(r[cols[1]])
                if v:
                    out.append({"date": d, "series": series, "place": str(r[cols[0]]).strip(), "brl_bag": v})
    return out


def fetch_brazil(full=False):
    rows = []
    for slug, name in NA_SERIES.items():
        dates = [None]
        if full:
            d0 = date(2018, 1, 1); d = date.today()
            dates = []
            while d > d0:
                dates.append(d); d -= timedelta(days=13)
        for d in dates:
            u = NA + slug + (f"/{d:%Y-%m-%d}" if d else "")
            r = get(u)
            if r is None or not r.ok:
                continue
            rows += parse_na(r.text, name)
            if d:
                time.sleep(0.4)
        log("NA", name, len(rows))
    if rows:
        merge_save(pd.DataFrame(rows), "br_physical.csv", ["date", "series", "place"])


# ---------------------------------------------------------------- Colombia (FNC)
def fetch_colombia(full=False):
    r = get("https://federaciondecafeteros.org/wp/estadisticas-cafeteras/")
    if r is None or not r.ok:
        log("fnc page fail"); return
    links = sorted(set(re.findall(r'href="([^"]+\.xlsx)"', r.text)))
    price = sorted([l for l in links if "Precios" in l], key=lambda u: re.search(r"uploads/(\d{4}/\d{2})", u).group(1))
    expo = sorted([l for l in links if "Exportaciones" in l], key=lambda u: re.search(r"uploads/(\d{4}/\d{2})", u).group(1))
    if price:
        x = get(price[-1], timeout=120)
        b = io.BytesIO(x.content)

        def sheet(name, cols, skip):
            d = pd.read_excel(b, sheet_name=name, header=None)
            d = d.iloc[:, cols].dropna(how="any")
            d.columns = skip
            d = d[pd.to_datetime(d.iloc[:, 0], errors="coerce").notna()]
            d.iloc[:, 0] = pd.to_datetime(d.iloc[:, 0]).dt.strftime("%Y-%m-%d")
            return d
        try:
            x1 = pd.ExcelFile(b)
            names = {n.split(".", 1)[1].strip().lower() if "." in n else n.lower(): n for n in x1.sheet_names}
            pick = lambda key: next(v for k, v in names.items() if key in k)
            dly = sheet(pick("interno diario"), [1, 2], ["date", "cop_125kg"])
            dly.to_csv(os.path.join(DATA, "col_internal_daily.csv"), index=False)
            ex = sheet(pick("ex_dock mensual"), [3, 4], ["month", "exdock_c_lb"])
            prod = sheet(pick("producción mensual") if any("producción mensual" in k for k in names) else pick("produc"), [3, 4], ["month", "prod_k_bags"])
            oic = pd.read_excel(b, sheet_name=pick("oic mensual"), header=None)
            hdr = oic.index[oic.iloc[:, 1].astype(str).str.strip() == "Mes"][0]
            oic = oic.iloc[hdr + 1:, 1:14]
            oic.columns = ["month", "composite", "col_ny", "col_eu", "col_w", "oth_ny", "oth_eu", "oth_w", "bra_ny", "bra_eu", "bra_w", "rob_ny", "rob_eu", "rob_w"][:oic.shape[1]]
            oic = oic[pd.to_datetime(oic.month, errors="coerce").notna()]
            oic["month"] = pd.to_datetime(oic.month).dt.strftime("%Y-%m")
            m = ex.assign(month=pd.to_datetime(ex.month).dt.strftime("%Y-%m")).merge(
                prod.assign(month=pd.to_datetime(prod.month).dt.strftime("%Y-%m")), on="month", how="outer")
            m.sort_values("month").to_csv(os.path.join(DATA, "col_monthly.csv"), index=False)
            oic.to_csv(os.path.join(DATA, "ico_groups_monthly.csv"), index=False)
            log("fnc prices", len(dly), len(m), len(oic))
        except Exception:
            traceback.print_exc()
    if expo:
        try:
            x = get(expo[-1], timeout=180)
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(x.content), read_only=True, data_only=True)
            ws = wb[[n for n in wb.sheetnames if "Total_Volumen" in n or n.startswith("1.")][0]]
            rows = []
            for row in ws.iter_rows(values_only=True):
                vals = [v for v in row if v is not None]
                if len(vals) >= 2 and isinstance(vals[0], datetime):
                    rows.append({"month": vals[0].strftime("%Y-%m"), "exports_k_bags": float(vals[1])})
            pd.DataFrame(rows).to_csv(os.path.join(DATA, "col_exports.csv"), index=False)
            log("fnc exports", len(rows))
        except Exception:
            traceback.print_exc()


# ---------------------------------------------------------------- Vietnam (giacaphe.com)
def _vn_day(d):
    u = "https://giacaphe.com/gia-ca-phe-noi-dia/" if d is None else f"https://giacaphe.com/gia-ca-phe-noi-dia-ngay-{d:%Y-%m-%d}/"
    r = get(u)
    if r is None or not r.ok:
        return None
    r.encoding = "utf-8"
    t = r.text
    m = re.search(r'data-cur="([\d.]+)"', t)
    if not m:
        m2 = re.search(r"trung bình\s*([\d,\.]+)\s*vnđ/kg", t, re.I)
        val = float(m2.group(1).replace(",", "").replace(".", "")) if m2 else None
    else:
        val = float(m.group(1))
    dm = re.search(r"ngày\s*(\d{2}/\d{2}/\d{4})", t, re.I)
    dd = datetime.strptime(dm.group(1), "%d/%m/%Y").date() if dm else (d or date.today())
    return {"date": dd.strftime("%Y-%m-%d"), "vnd_kg": val} if val else None


def fetch_vietnam(full=False):
    rows = []
    days = [None] + [date.today() - timedelta(days=k) for k in range(1, 10)]
    if full:
        d, d0 = date.today(), date(2021, 1, 1)
        while d > d0:
            days.append(d); d -= timedelta(days=7)
    for d in days:
        x = _vn_day(d)
        if x:
            rows.append(x)
        if d is not None:
            time.sleep(0.3)
    log("vn", len(rows))
    if rows:
        merge_save(pd.DataFrame(rows), "vn_physical.csv", ["date"])


# ---------------------------------------------------------------- ICO (coffee-analytics repo), PSD, COT
RAW = "https://raw.githubusercontent.com/BMR-Com/"


def fetch_ico():
    for f in ["prices", "exports", "certified_stocks"]:
        r = get(RAW + f"coffee-analytics/main/data/csv/{f}.csv")
        if r is not None and r.ok:
            open(os.path.join(DATA, f"ico_{f}.csv"), "wb").write(r.content); log("ico", f, len(r.content))


def fetch_psd():
    r = get("https://apps.fas.usda.gov/psdonline/downloads/psd_coffee_csv.zip", timeout=180)
    if r is None or not r.ok:
        return
    z = zipfile.ZipFile(io.BytesIO(r.content))
    d = pd.read_csv(io.BytesIO(z.read(z.namelist()[0])))
    d = d[d.Market_Year >= 2000][["Country_Code", "Country_Name", "Market_Year", "Attribute_Description", "Value"]]
    d.to_csv(os.path.join(DATA, "psd_coffee.csv"), index=False); log("psd", d.shape)


def fetch_cot():
    r = get(RAW + "cftc-cot-full/main/data/bcom_cot_master.csv", timeout=240)
    if r is None or not r.ok:
        log("cot fail"); return
    d = pd.read_csv(io.BytesIO(r.content))
    d = d[(d["commodity"] == "Coffee") & (d["crop_type"].astype(str) == "all")]
    keep = [c for c in d.columns if c in ("date", "commodity", "open_interest") or c.startswith(("managed_money", "prod_merc", "swap"))]
    d[keep].to_csv(os.path.join(DATA, "cot_coffee.csv"), index=False); log("cot", len(d))


if __name__ == "__main__":
    ap = argparse.ArgumentParser(); ap.add_argument("--full", action="store_true"); ap.add_argument("--only", default="")
    a = ap.parse_args()
    steps = {"futures": lambda: fetch_futures(a.full), "brazil": lambda: fetch_brazil(a.full), "colombia": lambda: fetch_colombia(a.full),
             "vietnam": lambda: fetch_vietnam(a.full), "ico": fetch_ico, "psd": fetch_psd, "cot": fetch_cot}
    for k, f in steps.items():
        if a.only and k not in a.only.split(","):
            continue
        try:
            f()
        except Exception:
            traceback.print_exc()
